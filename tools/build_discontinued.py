"""Rebuild Herby's discontinued list from the PlantMed inventory spreadsheet.

Usage:
    python3 tools/build_discontinued.py "<path to Inventory & Expiry .xlsx>" [herby.html]

Takes every product on the "Discontinued" tab, plus any row on "PLANTMED INVENTORY" whose
status is "Product DISCONTINUED", and writes discontinued.json. If a path to herby.html is
given, the list is also written into the page (const DISCONTINUED = [...]).
Herby still matches names against the pricelist; this list only labels discontinued scripts.
Needs openpyxl (pip install openpyxl).
"""
import json
import re
import sys

import openpyxl

wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
names = []


def add(name):
    name = re.sub(r"\s+", " ", str(name)).strip()
    if name and name.lower() not in {n.lower() for n in names}:
        names.append(name)


for row in wb["Discontinued"].iter_rows(min_row=4, values_only=True):
    if row[0] and not isinstance(row[0], (int, float)):
        add(row[0])

for row in wb["PLANTMED INVENTORY"].iter_rows(min_row=4, values_only=True):
    name, status = row[1], row[2]
    if name and status and "discontinued" in str(status).lower():
        add(name)

json.dump(names, open("discontinued.json", "w"), indent=0)
print(len(names), "discontinued products")

if len(sys.argv) > 2:
    page = open(sys.argv[2]).read()
    line = "const DISCONTINUED = " + json.dumps(names, separators=(",", ":")) + ";\n"
    if "const DISCONTINUED = " in page:
        page, n = re.subn(r"const DISCONTINUED = \[.*?\];\n", lambda m: line, page, count=1, flags=re.S)
    else:
        page, n = re.subn(r"(const CATALOG = \[.*?\];\n)", lambda m: m.group(1) + line, page, count=1, flags=re.S)
    assert n == 1, "could not place the list in the page"
    open(sys.argv[2], "w").write(page)
    print("updated", sys.argv[2])
